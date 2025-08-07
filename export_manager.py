"""
Multimedia Export and Sharing Manager
Handles various export formats and sharing capabilities
"""

import os
import json
import logging
import tempfile
import zipfile
import base64
from pathlib import Path
from typing import Dict, List, Any, Optional, Union
from dataclasses import dataclass, field
from datetime import datetime
import xml.etree.ElementTree as ET

import pandas as pd
from docx import Document
from docx.shared import Inches
import plotly.graph_objects as go
import plotly.express as px
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

from errors import FileProcessingError, ErrorCode

logger = logging.getLogger(__name__)


@dataclass
class ExportConfig:
    """Configuration for export operations"""
    format: str  # pdf, docx, json, csv, xml, html, txt
    include_metadata: bool = True
    include_timestamps: bool = True
    include_speaker_info: bool = True
    include_entities: bool = True
    include_insights: bool = True
    include_visualizations: bool = False
    custom_template: Optional[str] = None
    branding: Dict[str, Any] = field(default_factory=dict)


@dataclass
class ShareConfig:
    """Configuration for sharing operations"""
    platform: str  # email, link, cloud, api
    permissions: str = "read"  # read, write, admin
    expiry_hours: Optional[int] = None
    password_protected: bool = False
    notify_recipients: bool = True
    custom_message: Optional[str] = None


class MultimediaExporter:
    """Handles various export formats for transcription data"""
    
    def __init__(self, output_dir: str = "exports"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        
    def export_transcript(self, 
                         transcript_data: Dict[str, Any],
                         config: ExportConfig) -> str:
        """
        Export transcript data in specified format
        
        Args:
            transcript_data: Complete transcript data with metadata
            config: Export configuration
            
        Returns:
            str: Path to exported file
        """
        try:
            export_methods = {
                'pdf': self._export_pdf,
                'docx': self._export_docx,
                'json': self._export_json,
                'csv': self._export_csv,
                'xml': self._export_xml,
                'html': self._export_html,
                'txt': self._export_txt,
                'xlsx': self._export_xlsx,
                'md': self._export_markdown,
                'srt': self._export_srt,
                'vtt': self._export_vtt,
                'elan': self._export_elan
            }
            
            if config.format not in export_methods:
                raise FileProcessingError(
                    message=f"Unsupported export format: {config.format}",
                    error_code=ErrorCode.FILE_UNSUPPORTED_FORMAT,
                    user_message=f"Export format '{config.format}' is not supported."
                )
            
            export_method = export_methods[config.format]
            output_path = export_method(transcript_data, config)
            
            logger.info(f"Successfully exported to {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"Export failed: {str(e)}")
            if isinstance(e, FileProcessingError):
                raise
            raise FileProcessingError(
                message=f"Export operation failed: {str(e)}",
                error_code=ErrorCode.FILE_WRITE_ERROR,
                user_message="Failed to export transcript data."
            )
    
    def create_export_package(self, 
                            transcript_data: Dict[str, Any],
                            formats: List[str],
                            include_media: bool = False) -> str:
        """
        Create a comprehensive export package with multiple formats
        
        Args:
            transcript_data: Complete transcript data
            formats: List of export formats to include
            include_media: Whether to include original media files
            
        Returns:
            str: Path to ZIP package
        """
        try:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            package_name = f"transcript_export_{timestamp}.zip"
            package_path = self.output_dir / package_name
            
            with zipfile.ZipFile(package_path, 'w', zipfile.ZIP_DEFLATED) as zip_file:
                # Export in each requested format
                for format_type in formats:
                    config = ExportConfig(format=format_type)
                    export_path = self.export_transcript(transcript_data, config)
                    
                    # Add to ZIP with format-specific folder
                    arcname = f"{format_type}/{Path(export_path).name}"
                    zip_file.write(export_path, arcname)
                    
                    # Clean up individual export file
                    os.unlink(export_path)
                
                # Include visualizations if requested
                if any(fmt in formats for fmt in ['pdf', 'html']):
                    vis_files = self._create_visualizations(transcript_data)
                    for vis_file in vis_files:
                        arcname = f"visualizations/{Path(vis_file).name}"
                        zip_file.write(vis_file, arcname)
                        os.unlink(vis_file)
                
                # Include media files if requested
                if include_media and 'media_path' in transcript_data:
                    media_path = transcript_data['media_path']
                    if os.path.exists(media_path):
                        arcname = f"media/{Path(media_path).name}"
                        zip_file.write(media_path, arcname)
                
                # Add metadata file
                metadata = {
                    'export_date': datetime.now().isoformat(),
                    'formats_included': formats,
                    'transcript_id': transcript_data.get('id', 'unknown'),
                    'duration': transcript_data.get('duration', 0),
                    'speaker_count': len(transcript_data.get('speakers', [])),
                    'export_config': {
                        'include_media': include_media,
                        'formats': formats
                    }
                }
                
                metadata_json = json.dumps(metadata, indent=2)
                zip_file.writestr("export_metadata.json", metadata_json)
            
            logger.info(f"Created export package: {package_path}")
            return str(package_path)
            
        except Exception as e:
            logger.error(f"Package creation failed: {str(e)}")
            raise FileProcessingError(
                message=f"Failed to create export package: {str(e)}",
                error_code=ErrorCode.FILE_WRITE_ERROR,
                user_message="Failed to create export package."
            )
    
    def _export_pdf(self, data: Dict[str, Any], config: ExportConfig) -> str:
        """Export to PDF format"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"transcript_{timestamp}.pdf"
        output_path = self.output_dir / filename
        
        # Create PDF document
        doc = SimpleDocTemplate(str(output_path), pagesize=letter)
        styles = getSampleStyleSheet()
        story = []
        
        # Title
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            spaceAfter=30,
            textColor=colors.darkblue
        )
        
        title = data.get('title', 'Transcript Report')
        story.append(Paragraph(title, title_style))
        story.append(Spacer(1, 12))
        
        # Metadata section
        if config.include_metadata:
            story.append(Paragraph("Document Information", styles['Heading2']))
            
            metadata_data = [
                ['Generated', datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
                ['Duration', f"{data.get('duration', 0):.1f} seconds"],
                ['Language', data.get('language', 'Unknown')],
                ['Speakers', str(len(data.get('speakers', [])))]
            ]
            
            if 'confidence' in data:
                metadata_data.append(['Confidence', f"{data['confidence']:.1%}"])
            
            metadata_table = Table(metadata_data, colWidths=[2*inch, 3*inch])
            metadata_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            
            story.append(metadata_table)
            story.append(Spacer(1, 12))
        
        # Transcript content
        story.append(Paragraph("Transcript", styles['Heading2']))
        
        transcript_text = data.get('transcript', 'No transcript available')
        
        # Split into paragraphs for better formatting
        paragraphs = transcript_text.split('\n\n')
        for para in paragraphs:
            if para.strip():
                story.append(Paragraph(para.strip(), styles['Normal']))
                story.append(Spacer(1, 6))
        
        # Entities section
        if config.include_entities and 'entities' in data:
            story.append(Spacer(1, 12))
            story.append(Paragraph("Extracted Entities", styles['Heading2']))
            
            entities_by_type = {}
            for entity in data['entities']:
                entity_type = entity.get('label', 'Unknown')
                if entity_type not in entities_by_type:
                    entities_by_type[entity_type] = []
                entities_by_type[entity_type].append(entity.get('text', ''))
            
            for entity_type, entity_list in entities_by_type.items():
                story.append(Paragraph(f"{entity_type}:", styles['Heading3']))
                entity_text = ", ".join(set(entity_list))  # Remove duplicates
                story.append(Paragraph(entity_text, styles['Normal']))
                story.append(Spacer(1, 6))
        
        # Build PDF
        doc.build(story)
        return str(output_path)
    
    def _export_docx(self, data: Dict[str, Any], config: ExportConfig) -> str:
        """Export to DOCX format"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"transcript_{timestamp}.docx"
        output_path = self.output_dir / filename
        
        doc = Document()
        
        # Add title
        title = doc.add_heading(data.get('title', 'Transcript Report'), 0)
        
        # Add metadata
        if config.include_metadata:
            doc.add_heading('Document Information', level=1)
            
            metadata_table = doc.add_table(rows=1, cols=2)
            metadata_table.style = 'Table Grid'
            
            # Header row
            hdr_cells = metadata_table.rows[0].cells
            hdr_cells[0].text = 'Property'
            hdr_cells[1].text = 'Value'
            
            # Add metadata rows
            metadata_items = [
                ('Generated', datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
                ('Duration', f"{data.get('duration', 0):.1f} seconds"),
                ('Language', data.get('language', 'Unknown')),
                ('Speakers', str(len(data.get('speakers', []))))
            ]
            
            for prop, value in metadata_items:
                row_cells = metadata_table.add_row().cells
                row_cells[0].text = prop
                row_cells[1].text = value
        
        # Add transcript
        doc.add_heading('Transcript', level=1)
        transcript_text = data.get('transcript', 'No transcript available')
        doc.add_paragraph(transcript_text)
        
        # Add entities
        if config.include_entities and 'entities' in data:
            doc.add_heading('Extracted Entities', level=1)
            
            entities_by_type = {}
            for entity in data['entities']:
                entity_type = entity.get('label', 'Unknown')
                if entity_type not in entities_by_type:
                    entities_by_type[entity_type] = []
                entities_by_type[entity_type].append(entity.get('text', ''))
            
            for entity_type, entity_list in entities_by_type.items():
                doc.add_heading(entity_type, level=2)
                entity_text = ", ".join(set(entity_list))
                doc.add_paragraph(entity_text)
        
        doc.save(str(output_path))
        return str(output_path)
    
    def _export_json(self, data: Dict[str, Any], config: ExportConfig) -> str:
        """Export to JSON format"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"transcript_{timestamp}.json"
        output_path = self.output_dir / filename
        
        # Prepare export data
        export_data = {
            'export_info': {
                'generated_at': datetime.now().isoformat(),
                'format': 'json',
                'version': '1.0'
            }
        }
        
        # Add requested sections
        if config.include_metadata:
            export_data['metadata'] = {
                'id': data.get('id'),
                'title': data.get('title'),
                'duration': data.get('duration'),
                'language': data.get('language'),
                'confidence': data.get('confidence'),
                'created_at': data.get('created_at'),
                'speakers': data.get('speakers', [])
            }
        
        export_data['transcript'] = data.get('transcript', '')
        
        if config.include_entities and 'entities' in data:
            export_data['entities'] = data['entities']
        
        if config.include_insights and 'insights' in data:
            export_data['insights'] = data['insights']
        
        if config.include_speaker_info and 'speaker_segments' in data:
            export_data['speaker_segments'] = data['speaker_segments']
        
        # Write JSON file
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(export_data, f, indent=2, ensure_ascii=False)
        
        return str(output_path)
    
    def _export_csv(self, data: Dict[str, Any], config: ExportConfig) -> str:
        """Export to CSV format"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"transcript_{timestamp}.csv"
        output_path = self.output_dir / filename
        
        # Prepare CSV data
        rows = []
        
        # If speaker segments available, use them
        if config.include_speaker_info and 'speaker_segments' in data:
            for segment in data['speaker_segments']:
                row = {
                    'timestamp': segment.get('start_time', 0),
                    'duration': segment.get('duration', 0),
                    'speaker': segment.get('speaker_id', 'Unknown'),
                    'text': segment.get('text', ''),
                    'confidence': segment.get('confidence', 0)
                }
                rows.append(row)
        else:
            # Single row with full transcript
            row = {
                'timestamp': 0,
                'duration': data.get('duration', 0),
                'speaker': 'All',
                'text': data.get('transcript', ''),
                'confidence': data.get('confidence', 0)
            }
            rows.append(row)
        
        # Create DataFrame and save
        df = pd.DataFrame(rows)
        df.to_csv(output_path, index=False, encoding='utf-8')
        
        return str(output_path)
    
    def _export_xml(self, data: Dict[str, Any], config: ExportConfig) -> str:
        """Export to XML format"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"transcript_{timestamp}.xml"
        output_path = self.output_dir / filename
        
        # Create XML structure
        root = ET.Element("transcript")
        
        # Add metadata
        if config.include_metadata:
            metadata_elem = ET.SubElement(root, "metadata")
            
            metadata_fields = {
                'id': data.get('id'),
                'title': data.get('title'),
                'duration': str(data.get('duration', 0)),
                'language': data.get('language'),
                'confidence': str(data.get('confidence', 0)),
                'generated_at': datetime.now().isoformat()
            }
            
            for key, value in metadata_fields.items():
                if value:
                    elem = ET.SubElement(metadata_elem, key)
                    elem.text = value
        
        # Add transcript content
        content_elem = ET.SubElement(root, "content")
        content_elem.text = data.get('transcript', '')
        
        # Add entities
        if config.include_entities and 'entities' in data:
            entities_elem = ET.SubElement(root, "entities")
            
            for entity in data['entities']:
                entity_elem = ET.SubElement(entities_elem, "entity")
                entity_elem.set("type", entity.get('label', 'Unknown'))
                entity_elem.set("confidence", str(entity.get('confidence', 0)))
                entity_elem.text = entity.get('text', '')
        
        # Add speaker segments
        if config.include_speaker_info and 'speaker_segments' in data:
            segments_elem = ET.SubElement(root, "speaker_segments")
            
            for segment in data['speaker_segments']:
                segment_elem = ET.SubElement(segments_elem, "segment")
                segment_elem.set("speaker", segment.get('speaker_id', 'Unknown'))
                segment_elem.set("start_time", str(segment.get('start_time', 0)))
                segment_elem.set("duration", str(segment.get('duration', 0)))
                segment_elem.text = segment.get('text', '')
        
        # Write XML file
        tree = ET.ElementTree(root)
        tree.write(output_path, encoding='utf-8', xml_declaration=True)
        
        return str(output_path)
    
    def _export_html(self, data: Dict[str, Any], config: ExportConfig) -> str:
        """Export to HTML format"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"transcript_{timestamp}.html"
        output_path = self.output_dir / filename
        
        # HTML template
        html_content = f"""
        <!DOCTYPE html>
        <html lang="en">
        <head>
            <meta charset="UTF-8">
            <meta name="viewport" content="width=device-width, initial-scale=1.0">
            <title>{data.get('title', 'Transcript Report')}</title>
            <style>
                body {{ font-family: Arial, sans-serif; margin: 40px; line-height: 1.6; }}
                .header {{ background-color: #f4f4f4; padding: 20px; border-radius: 5px; margin-bottom: 30px; }}
                .metadata {{ background-color: #e9e9e9; padding: 15px; border-radius: 5px; margin-bottom: 20px; }}
                .transcript {{ background-color: #fff; padding: 20px; border: 1px solid #ddd; border-radius: 5px; }}
                .entities {{ margin-top: 30px; }}
                .entity-type {{ background-color: #f0f8ff; padding: 10px; margin: 10px 0; border-radius: 3px; }}
                .timestamp {{ color: #666; font-size: 0.9em; }}
                h1 {{ color: #333; }}
                h2 {{ color: #555; border-bottom: 2px solid #eee; padding-bottom: 10px; }}
                h3 {{ color: #777; }}
            </style>
        </head>
        <body>
            <div class="header">
                <h1>{data.get('title', 'Transcript Report')}</h1>
                <p class="timestamp">Generated on {datetime.now().strftime('%Y-%m-%d at %H:%M:%S')}</p>
            </div>
        """
        
        # Add metadata
        if config.include_metadata:
            html_content += """
            <div class="metadata">
                <h2>Document Information</h2>
                <ul>
            """
            
            metadata_items = [
                ('Duration', f"{data.get('duration', 0):.1f} seconds"),
                ('Language', data.get('language', 'Unknown')),
                ('Speakers', str(len(data.get('speakers', [])))),
                ('Confidence', f"{data.get('confidence', 0):.1%}" if 'confidence' in data else 'N/A')
            ]
            
            for prop, value in metadata_items:
                html_content += f"<li><strong>{prop}:</strong> {value}</li>"
            
            html_content += "</ul></div>"
        
        # Add transcript
        html_content += f"""
            <div class="transcript">
                <h2>Transcript</h2>
                <p>{data.get('transcript', 'No transcript available').replace(chr(10), '<br>')}</p>
            </div>
        """
        
        # Add entities
        if config.include_entities and 'entities' in data:
            html_content += '<div class="entities"><h2>Extracted Entities</h2>'
            
            entities_by_type = {}
            for entity in data['entities']:
                entity_type = entity.get('label', 'Unknown')
                if entity_type not in entities_by_type:
                    entities_by_type[entity_type] = []
                entities_by_type[entity_type].append(entity.get('text', ''))
            
            for entity_type, entity_list in entities_by_type.items():
                html_content += f"""
                <div class="entity-type">
                    <h3>{entity_type}</h3>
                    <p>{', '.join(set(entity_list))}</p>
                </div>
                """
            
            html_content += '</div>'
        
        html_content += """
        </body>
        </html>
        """
        
        # Write HTML file
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(html_content)
        
        return str(output_path)
    
    def _export_txt(self, data: Dict[str, Any], config: ExportConfig) -> str:
        """Export to plain text format"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"transcript_{timestamp}.txt"
        output_path = self.output_dir / filename
        
        content = []
        
        # Header
        content.append("TRANSCRIPT REPORT")
        content.append("=" * 50)
        content.append(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        content.append("")
        
        # Metadata
        if config.include_metadata:
            content.append("DOCUMENT INFORMATION")
            content.append("-" * 20)
            content.append(f"Title: {data.get('title', 'Untitled')}")
            content.append(f"Duration: {data.get('duration', 0):.1f} seconds")
            content.append(f"Language: {data.get('language', 'Unknown')}")
            content.append(f"Speakers: {len(data.get('speakers', []))}")
            if 'confidence' in data:
                content.append(f"Confidence: {data['confidence']:.1%}")
            content.append("")
        
        # Transcript
        content.append("TRANSCRIPT")
        content.append("-" * 10)
        content.append(data.get('transcript', 'No transcript available'))
        content.append("")
        
        # Entities
        if config.include_entities and 'entities' in data:
            content.append("EXTRACTED ENTITIES")
            content.append("-" * 18)
            
            entities_by_type = {}
            for entity in data['entities']:
                entity_type = entity.get('label', 'Unknown')
                if entity_type not in entities_by_type:
                    entities_by_type[entity_type] = []
                entities_by_type[entity_type].append(entity.get('text', ''))
            
            for entity_type, entity_list in entities_by_type.items():
                content.append(f"{entity_type}: {', '.join(set(entity_list))}")
            
            content.append("")
        
        # Write text file
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(content))
        
        return str(output_path)
    
    def _export_xlsx(self, data: Dict[str, Any], config: ExportConfig) -> str:
        """Export to Excel format"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"transcript_{timestamp}.xlsx"
        output_path = self.output_dir / filename
        
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            # Transcript sheet
            if config.include_speaker_info and 'speaker_segments' in data:
                segments_data = []
                for segment in data['speaker_segments']:
                    segments_data.append({
                        'Start Time': segment.get('start_time', 0),
                        'Duration': segment.get('duration', 0),
                        'Speaker': segment.get('speaker_id', 'Unknown'),
                        'Text': segment.get('text', ''),
                        'Confidence': segment.get('confidence', 0)
                    })
                
                segments_df = pd.DataFrame(segments_data)
                segments_df.to_excel(writer, sheet_name='Transcript', index=False)
            
            # Entities sheet
            if config.include_entities and 'entities' in data:
                entities_data = []
                for entity in data['entities']:
                    entities_data.append({
                        'Entity': entity.get('text', ''),
                        'Type': entity.get('label', 'Unknown'),
                        'Confidence': entity.get('confidence', 0),
                        'Start Position': entity.get('start', 0),
                        'End Position': entity.get('end', 0)
                    })
                
                entities_df = pd.DataFrame(entities_data)
                entities_df.to_excel(writer, sheet_name='Entities', index=False)
            
            # Metadata sheet
            if config.include_metadata:
                metadata_data = {
                    'Property': ['Title', 'Duration', 'Language', 'Speakers', 'Generated'],
                    'Value': [
                        data.get('title', 'Untitled'),
                        f"{data.get('duration', 0):.1f} seconds",
                        data.get('language', 'Unknown'),
                        len(data.get('speakers', [])),
                        datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                    ]
                }
                
                metadata_df = pd.DataFrame(metadata_data)
                metadata_df.to_excel(writer, sheet_name='Metadata', index=False)
        
        return str(output_path)
    
    def _export_markdown(self, data: Dict[str, Any], config: ExportConfig) -> str:
        """Export to Markdown format"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"transcript_{timestamp}.md"
        output_path = self.output_dir / filename
        
        content = []
        
        # Header
        content.append(f"# {data.get('title', 'Transcript Report')}")
        content.append("")
        content.append(f"*Generated on {datetime.now().strftime('%Y-%m-%d at %H:%M:%S')}*")
        content.append("")
        
        # Metadata
        if config.include_metadata:
            content.append("## Document Information")
            content.append("")
            content.append(f"- **Duration:** {data.get('duration', 0):.1f} seconds")
            content.append(f"- **Language:** {data.get('language', 'Unknown')}")
            content.append(f"- **Speakers:** {len(data.get('speakers', []))}")
            if 'confidence' in data:
                content.append(f"- **Confidence:** {data['confidence']:.1%}")
            content.append("")
        
        # Transcript
        content.append("## Transcript")
        content.append("")
        content.append(data.get('transcript', 'No transcript available'))
        content.append("")
        
        # Entities
        if config.include_entities and 'entities' in data:
            content.append("## Extracted Entities")
            content.append("")
            
            entities_by_type = {}
            for entity in data['entities']:
                entity_type = entity.get('label', 'Unknown')
                if entity_type not in entities_by_type:
                    entities_by_type[entity_type] = []
                entities_by_type[entity_type].append(entity.get('text', ''))
            
            for entity_type, entity_list in entities_by_type.items():
                content.append(f"### {entity_type}")
                content.append("")
                for entity in set(entity_list):
                    content.append(f"- {entity}")
                content.append("")
        
        # Write markdown file
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(content))
        
        return str(output_path)
    
    def _create_visualizations(self, data: Dict[str, Any]) -> List[str]:
        """Create visualization files for the transcript data"""
        vis_files = []
        
        try:
            # Speaker distribution chart
            if 'speakers' in data and len(data['speakers']) > 1:
                fig = px.pie(
                    values=[1] * len(data['speakers']),
                    names=data['speakers'],
                    title="Speaker Distribution"
                )
                
                vis_path = self.output_dir / f"speaker_distribution_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html"
                fig.write_html(str(vis_path))
                vis_files.append(str(vis_path))
            
            # Entity types chart
            if 'entities' in data:
                entity_counts = {}
                for entity in data['entities']:
                    entity_type = entity.get('label', 'Unknown')
                    entity_counts[entity_type] = entity_counts.get(entity_type, 0) + 1
                
                if entity_counts:
                    fig = px.bar(
                        x=list(entity_counts.keys()),
                        y=list(entity_counts.values()),
                        title="Entity Types Distribution"
                    )
                    
                    vis_path = self.output_dir / f"entity_distribution_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html"
                    fig.write_html(str(vis_path))
                    vis_files.append(str(vis_path))
        
        except Exception as e:
            logger.warning(f"Failed to create visualizations: {str(e)}")
        
        return vis_files


class SharingManager:
    """Handles sharing and distribution of exported content"""
    
    def __init__(self):
        pass
    
    def create_shareable_link(self, 
                           file_path: str,
                           config: ShareConfig) -> Dict[str, Any]:
        """
        Create a shareable link for exported content
        
        Args:
            file_path: Path to file to share
            config: Sharing configuration
            
        Returns:
            Dict with sharing information
        """
        try:
            # Generate unique share ID
            import uuid
            share_id = str(uuid.uuid4())
            
            # Create share metadata
            share_info = {
                'share_id': share_id,
                'file_path': file_path,
                'created_at': datetime.now().isoformat(),
                'expires_at': None,
                'permissions': config.permissions,
                'platform': config.platform,
                'password_protected': config.password_protected,
                'access_count': 0,
                'shareable_link': f"https://example.com/share/{share_id}"
            }
            
            if config.expiry_hours:
                from datetime import timedelta
                expires_at = datetime.now() + timedelta(hours=config.expiry_hours)
                share_info['expires_at'] = expires_at.isoformat()
            
            logger.info(f"Created shareable link: {share_id}")
            return share_info
            
        except Exception as e:
            logger.error(f"Failed to create shareable link: {str(e)}")
            raise FileProcessingError(
                message=f"Failed to create shareable link: {str(e)}",
                error_code=ErrorCode.PROCESSING_ERROR,
                user_message="Failed to create shareable link."
            )
    
    def generate_email_content(self, 
                             share_info: Dict[str, Any],
                             custom_message: str = None) -> Dict[str, str]:
        """Generate email content for sharing"""
        subject = f"Transcript Shared: {share_info.get('title', 'Document')}"
        
        body = f"""
Hello,

A transcript document has been shared with you.

{custom_message or 'Please find the shared document using the link below.'}

Access Link: {share_info['shareable_link']}
Permissions: {share_info['permissions']}

"""
        
        if share_info.get('expires_at'):
            body += f"This link expires on: {share_info['expires_at']}\n"
        
        if share_info.get('password_protected'):
            body += "This document is password protected. Please contact the sender for the password.\n"
        
        body += """
Best regards,
Transcript System
        """
        
        return {
            'subject': subject,
            'body': body
        }
    
    def _export_srt(self, data: Dict[str, Any], config: ExportConfig) -> str:
        """Export to SRT subtitle format"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"transcript_{timestamp}.srt"
        output_path = self.output_dir / filename
        
        srt_content = []
        segments = data.get('segments', [])
        
        if not segments:
            # Create single subtitle from full transcript
            segments = [{
                'start_time': 0,
                'end_time': data.get('duration', 60),
                'text': data.get('transcript', ''),
                'speaker': 'Speaker'
            }]
        
        for i, segment in enumerate(segments, 1):
            start_time = self._format_srt_time(segment.get('start_time', 0))
            end_time = self._format_srt_time(segment.get('end_time', 0))
            text = segment.get('text', '').strip()
            
            if config.include_speaker_info and segment.get('speaker'):
                text = f"[{segment['speaker']}] {text}"
            
            srt_entry = f"{i}\n{start_time} --> {end_time}\n{text}\n\n"
            srt_content.append(srt_entry)
        
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(''.join(srt_content))
        
        logger.info(f"SRT subtitle file created: {output_path}")
        return str(output_path)
    
    def _export_vtt(self, data: Dict[str, Any], config: ExportConfig) -> str:
        """Export to WebVTT subtitle format"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"transcript_{timestamp}.vtt"
        output_path = self.output_dir / filename
        
        vtt_content = ["WEBVTT\n\n"]
        segments = data.get('segments', [])
        
        if not segments:
            # Create single subtitle from full transcript
            segments = [{
                'start_time': 0,
                'end_time': data.get('duration', 60),
                'text': data.get('transcript', ''),
                'speaker': 'Speaker'
            }]
        
        for segment in segments:
            start_time = self._format_vtt_time(segment.get('start_time', 0))
            end_time = self._format_vtt_time(segment.get('end_time', 0))
            text = segment.get('text', '').strip()
            
            if config.include_speaker_info and segment.get('speaker'):
                text = f"<v {segment['speaker']}>{text}"
            
            vtt_entry = f"{start_time} --> {end_time}\n{text}\n\n"
            vtt_content.append(vtt_entry)
        
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(''.join(vtt_content))
        
        logger.info(f"VTT subtitle file created: {output_path}")
        return str(output_path)
    
    def _export_elan(self, data: Dict[str, Any], config: ExportConfig) -> str:
        """Export to ELAN annotation format"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"transcript_{timestamp}.eaf"
        output_path = self.output_dir / filename
        
        # Create ELAN XML structure
        root = ET.Element("ANNOTATION_DOCUMENT", {
            "AUTHOR": "Transcript System",
            "DATE": datetime.now().isoformat(),
            "FORMAT": "3.0",
            "VERSION": "3.0"
        })
        
        # Header
        header = ET.SubElement(root, "HEADER", {
            "MEDIA_FILE": "",
            "TIME_UNITS": "milliseconds"
        })
        
        # Media descriptor
        if data.get('media_file'):
            media_desc = ET.SubElement(header, "MEDIA_DESCRIPTOR", {
                "MEDIA_URL": data['media_file'],
                "MIME_TYPE": data.get('media_type', 'audio/wav')
            })
        
        # Time order
        time_order = ET.SubElement(root, "TIME_ORDER")
        
        # Create time slots
        segments = data.get('segments', [])
        time_slot_id = 1
        time_slots = {}
        
        for segment in segments:
            start_time = int(segment.get('start_time', 0) * 1000)  # Convert to milliseconds
            end_time = int(segment.get('end_time', 0) * 1000)
            
            start_slot = f"ts{time_slot_id}"
            end_slot = f"ts{time_slot_id + 1}"
            
            ET.SubElement(time_order, "TIME_SLOT", {
                "TIME_SLOT_ID": start_slot,
                "TIME_VALUE": str(start_time)
            })
            
            ET.SubElement(time_order, "TIME_SLOT", {
                "TIME_SLOT_ID": end_slot,
                "TIME_VALUE": str(end_time)
            })
            
            time_slots[segment.get('id', len(time_slots))] = (start_slot, end_slot)
            time_slot_id += 2
        
        # Create tier
        tier = ET.SubElement(root, "TIER", {
            "LINGUISTIC_TYPE_REF": "default-lt",
            "TIER_ID": "transcript"
        })
        
        # Add annotations
        for i, segment in enumerate(segments):
            annotation = ET.SubElement(tier, "ANNOTATION")
            alignable_annotation = ET.SubElement(annotation, "ALIGNABLE_ANNOTATION", {
                "ANNOTATION_ID": f"a{i+1}",
                "TIME_SLOT_REF1": time_slots[segment.get('id', i)][0],
                "TIME_SLOT_REF2": time_slots[segment.get('id', i)][1]
            })
            
            annotation_value = ET.SubElement(alignable_annotation, "ANNOTATION_VALUE")
            text = segment.get('text', '').strip()
            
            if config.include_speaker_info and segment.get('speaker'):
                text = f"[{segment['speaker']}] {text}"
            
            annotation_value.text = text
        
        # Linguistic type
        linguistic_type = ET.SubElement(root, "LINGUISTIC_TYPE", {
            "GRAPHIC_REFERENCES": "false",
            "LINGUISTIC_TYPE_ID": "default-lt",
            "TIME_ALIGNABLE": "true"
        })
        
        # Write to file
        tree = ET.ElementTree(root)
        ET.indent(tree, space="  ", level=0)
        tree.write(output_path, encoding='utf-8', xml_declaration=True)
        
        logger.info(f"ELAN annotation file created: {output_path}")
        return str(output_path)
    
    def _format_srt_time(self, seconds: float) -> str:
        """Format time for SRT format (HH:MM:SS,mmm)"""
        hours = int(seconds // 3600)
        minutes = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        millisecs = int((seconds % 1) * 1000)
        return f"{hours:02d}:{minutes:02d}:{secs:02d},{millisecs:03d}"
    
    def _format_vtt_time(self, seconds: float) -> str:
        """Format time for VTT format (HH:MM:SS.mmm)"""
        hours = int(seconds // 3600)
        minutes = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        millisecs = int((seconds % 1) * 1000)
        return f"{hours:02d}:{minutes:02d}:{secs:02d}.{millisecs:03d}"